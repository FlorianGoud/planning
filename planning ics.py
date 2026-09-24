#!/usr/bin/env python3
"""
Planning ITECH 1 (Excel) + fichier des salles (Excel)  ->  calendrier .ics

Usage :
    python planning_ics.py --download --tous --outdir docs      # G1, G2, G3 + sous-groupes A/B (9 fichiers)
    python planning_ics.py PLANNING.xlsx SALLES.xlsx --tous      # ou à partir de fichiers locaux
    python planning_ics.py PLANNING.xlsx SALLES.xlsx --groupe G2 -o itech1_G2.ics
    python planning_ics.py PLANNING.xlsx SALLES.xlsx --groupe G2 --sous-groupe B -o itech1_G2B.ics

Dépendance : pip install openpyxl
"""
import argparse, datetime as dt, hashlib, os, re, unicodedata, urllib.request
from collections import defaultdict
from openpyxl import load_workbook

# ----------------------------------------------------------------- réglages
FIRST_MONDAY = dt.date(2026, 8, 24)   # lundi de la semaine 35 (la cellule C2 du fichier contient 2024 par erreur)
FIRST_COL = 3                          # colonne C : 1re colonne de la semaine 35
COLS_PER_WEEK = 9                      # 3 groupes x 3 colonnes
FIRST_ROW = 5                          # ligne du Lundi 08h00
ROWS_PER_DAY = 21                      # 08h00 -> 18h30 par demi-heures (12h-13h = 2 lignes)
DAY_START_MIN = 8 * 60
GROUP_COLS = {"G1": (0, 2), "G2": (3, 5), "G3": (6, 8)}
ROOM_COLS = range(3, 28)               # C..AA dans "Salles de cours"
TZ = "Europe/Paris"
PLANNING_ID = "1ZfgKO-68K1ioRlTjY2x6Z5C8sH1TpJ19nxcdbMMIyPI"   # Google Sheet du planning (lecture publique)
SALLES_ID = "1es1RFXRtJGeyL8MzFaw6fyhEzXDbumx6GTGYiUxBSAg"     # Google Sheet des salles (lecture publique)
MIN_EVENTS = 50                        # garde-fou : en dessous, on n'écrase rien

STOP = {"td", "tp", "ds", "cm", "de", "des", "du", "la", "le", "les", "l", "d", "a", "au", "et", "en",
        "itech", "1", "2", "3", "e", "learning", "initiation", "introduction"}
SYN = {"legislation": "droit", "travail": "droit", "economique": "eco", "economie": "eco",
       "resistance": "rdm", "materiaux": "rdm", "fluides": "flu", "mecanique": "meca",
       "thermodynamique": "thermo", "organique": "orga", "incertitudes": "incertitude",
       "analyses": "analyse", "instrumentales": "instrumentale", "scientifiques": "bsi",
       "bases": "bsi", "ingenieur": "bsi", "outils": "oin", "numeriques": "oin", "informatiques": "oin",
       "generale": "chimie", "profil": "profil", "parole": "parole"}


# ----------------------------------------------------------------- utilitaires texte
def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\s*&\s*", "&", s)
    return re.sub(r"[^a-z0-9&]+", " ", s).strip()


def tokens(title):
    out = set()
    for w in norm(title).split():
        if w in STOP or w.isdigit():
            continue
        out.add(SYN.get(w, w))
    return out


def tok_match(a, b):
    return any(x == y or (len(x) >= 3 and len(y) >= 3 and (x.startswith(y) or y.startswith(x)))
               for x in a for y in b)


TEACHER_RE = re.compile(r"^(?:[A-Z]{1,2}\.\s?|[A-Z]\s)([A-ZÀ-Ý][a-zà-ÿ'\-]{2,})")


def teachers(lines):
    """Noms de famille repérés dans des lignes type 'J.Champliaud', 'AC.Besson' ou 'L Scalone'."""
    out = set()
    for ln in lines:
        m = TEACHER_RE.match(ln.strip())
        if m:
            out.add(norm(m.group(1)))
    return out


def content_lines(lines):
    """Lignes qui décrivent le cours (ni prof, ni compteur '1/3', ni '- 1 -')."""
    return [l for l in lines if not TEACHER_RE.match(l) and not re.fullmatch(r"[-\s\d/]+", l)]


def group_marker(text):
    """Groupe et sous-groupe cités dans une réservation de salle : 'Grp 2A', 'G2 B', 'G3'..."""
    m = re.search(r"\b(?:g|grp|groupe)\s*([123])", norm(text))
    if not m:
        return None, None
    sub = re.search(r"\b(?:G|Grp|Groupe)\s*[123]\s?([AB])\b(?![.'’])", text)   # 'G3 A.Allès' n'est pas un sous-groupe
    return f"G{m.group(1)}", (sub.group(1) if sub else None)


# ----------------------------------------------------------------- lecture du planning
def week_starts(ws, n=60):
    """Lundi de chaque semaine, lu sur la ligne 2 (en ignorant les dates aberrantes, ex. 2024 en C2)."""
    raw = {}
    for k in range(n):
        v = ws.cell(2, FIRST_COL + COLS_PER_WEEK * k).value
        if isinstance(v, dt.datetime) and v.year >= 2025 and v.weekday() == 0:
            raw[k] = v.date()
    anchor = (min(raw.items())[1] - dt.timedelta(days=7 * min(raw))) if raw else FIRST_MONDAY
    return [raw.get(k, anchor + dt.timedelta(days=7 * k)) for k in range(n)]


def read_planning(path):
    ws = load_workbook(path, data_only=True).worksheets[0]
    starts = week_starts(ws)
    merged = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
    events = []
    for row in ws.iter_rows(min_row=FIRST_ROW, max_row=FIRST_ROW + 5 * ROWS_PER_DAY - 1, min_col=FIRST_COL):
        for c in row:
            if c.value is None or not str(c.value).strip():
                continue
            m = merged.get((c.row, c.column))
            r2, c2 = (m.max_row, m.max_col) if m else (c.row, c.column)
            week = (c.column - FIRST_COL) // COLS_PER_WEEK
            p1, p2 = (c.column - FIRST_COL) % COLS_PER_WEEK, (c2 - FIRST_COL) % COLS_PER_WEEK
            if c2 - c.column >= COLS_PER_WEEK:
                p2 = COLS_PER_WEEK - 1
            day, off = divmod(c.row - FIRST_ROW, ROWS_PER_DAY)
            n_rows = r2 - c.row + 1
            groups = [g for g, (a, b) in GROUP_COLS.items() if p1 <= b and p2 >= a]
            # sous-groupes : dans chaque groupe, A = 2 premières colonnes, B = 3e colonne
            subs = {}
            for g in groups:
                a, b = GROUP_COLS[g]
                lo, hi = max(p1, a) - a, min(p2, b) - a
                subs[g] = ({"A"} if lo <= 1 else set()) | ({"B"} if hi >= 2 else set())
                if len(groups) > 1:          # événement commun à plusieurs groupes : pas de sous-groupe
                    subs[g] = {"A", "B"}
            base = starts[week] + dt.timedelta(days=day)
            text = str(c.value).strip()
            if off == 0 and n_rows >= ROWS_PER_DAY:      # jour(s) entier(s) : vacances, férié...
                ndays = max(1, min(5 - day, round(n_rows / ROWS_PER_DAY)))
                events.append(dict(date=base, end_date=base + dt.timedelta(days=ndays), allday=True,
                                   text=text, groups=groups, subs=subs))
                continue
            r2 = min(r2, FIRST_ROW + (day + 1) * ROWS_PER_DAY - 1)
            start = DAY_START_MIN + 30 * off
            end = DAY_START_MIN + 30 * (r2 - (FIRST_ROW + day * ROWS_PER_DAY) + 1)
            # horaire écrit dans le texte (ex. "13h30 - 15h30") : il prime sur le bloc
            t = re.search(r"(\d{1,2})\s*h\s*(\d{2})?\s*-\s*(\d{1,2})\s*h\s*(\d{2})?", text)
            if t:
                s2 = int(t.group(1)) * 60 + int(t.group(2) or 0)
                e2 = int(t.group(3)) * 60 + int(t.group(4) or 0)
                if 6 * 60 <= s2 < e2 <= 22 * 60:
                    start, end = s2, e2
            # sous-groupe écrit dans le titre d'un cours d'un seul groupe (ex. "TP Biblio 2A")
            tm = re.search(r"(?<![A-Za-z0-9])([123])([AB])(?![A-Za-z0-9])", text.split("\n")[0])
            if tm and len(groups) == 1 and f"G{tm.group(1)}" == groups[0]:
                subs[groups[0]] = {tm.group(2)}
            events.append(dict(date=base, start=start, end=end, allday=False, text=text,
                               groups=groups, subs=subs))
    return events


# ----------------------------------------------------------------- lecture des salles
def read_rooms(path):
    ws = load_workbook(path, data_only=True)["Salles de cours"]
    hdr = {c: str(ws.cell(3, c).value or "").split("\n")[0].strip() for c in ROOM_COLS}
    merged = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
    out = []
    for r in range(4, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if not isinstance(d, dt.datetime):
            continue
        for rr in range(r, r + ROWS_PER_DAY):
            for c in ROOM_COLS:
                v = ws.cell(rr, c).value
                if v is None or not str(v).strip():
                    continue
                m = merged.get((rr, c))
                r2 = min(m.max_row if m else rr, r + ROWS_PER_DAY - 1)
                out.append(dict(date=d.date(), start=DAY_START_MIN + 30 * (rr - r),
                                end=DAY_START_MIN + 30 * (r2 - r + 1), room=hdr[c], text=str(v).strip()))
    return out


# ----------------------------------------------------------------- jointure planning <-> salles
def find_rooms(ev, group, rooms_by_date, sub=None):
    if ev["allday"]:
        return []
    lines = [l.strip() for l in ev["text"].split("\n") if l.strip()]
    if norm(lines[0]).startswith("e learning"):
        return []
    title_tok = set().union(*[tokens(l) for l in content_lines(lines)[:2]]) if content_lines(lines) else set()
    if not title_tok:
        return []
    ev_teachers = teachers(lines)
    found = []
    for r in rooms_by_date.get(ev["date"], []):
        rl = [l.strip() for l in r["text"].split("\n") if l.strip()]
        if not re.search(r"itech\s*1", norm(r["text"])) and not (ev_teachers & teachers(rl)):
            continue
        ov = min(ev["end"], r["end"]) - max(ev["start"], r["start"])
        if ov < 0.6 * min(ev["end"] - ev["start"], r["end"] - r["start"]):
            continue
        if not tok_match(title_tok, set().union(*[tokens(l) for l in content_lines(rl)[:1]])):
            continue
        mk, msub = group_marker(r["text"])
        if mk and group and mk != group:
            continue
        want = {sub} if sub else ev["subs"].get(group, set())
        if msub and len(want) == 1 and msub not in want:
            continue
        rt = teachers(rl)
        if ev_teachers and rt and not (ev_teachers & rt):
            continue
        found.append(dict(room=r["room"], teacher=", ".join(sorted(t.title() for t in rt)) if rt else "", marked=bool(mk)))
    # si certaines salles portent explicitement le bon groupe, on écarte les autres
    if any(f["marked"] for f in found):
        found = [f for f in found if f["marked"]]
    uniq, seen = [], set()
    for f in found:
        if f["room"] not in seen:
            seen.add(f["room"]); uniq.append(f)
    return uniq


# ----------------------------------------------------------------- écriture .ics
def esc(s):
    return s.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def fold(line):
    b, out = line.encode("utf-8"), []
    while len(b) > 74:
        cut = 74
        while (b[cut] & 0xC0) == 0x80:
            cut -= 1
        out.append(b[:cut].decode("utf-8")); b = b[cut:]
        b = b" " + b
    out.append(b.decode("utf-8"))
    return "\r\n".join(out)


VTZ = """BEGIN:VTIMEZONE
TZID:Europe/Paris
BEGIN:DAYLIGHT
TZOFFSETFROM:+0100
TZOFFSETTO:+0200
TZNAME:CEST
DTSTART:19700329T020000
RRULE:FREQ=YEARLY;BYMONTH=3;BYDAY=-1SU
END:DAYLIGHT
BEGIN:STANDARD
TZOFFSETFROM:+0200
TZOFFSETTO:+0100
TZNAME:CET
DTSTART:19701025T030000
RRULE:FREQ=YEARLY;BYMONTH=10;BYDAY=-1SU
END:STANDARD
END:VTIMEZONE""".replace("\n", "\r\n")


def build_ics(events, group, rooms_by_date, sub=None):
    label = group + (sub or "")
    stamp = "20260901T000000Z"   # fixe : le fichier ne change que si le planning change
    L = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//planning-itech//FR", "CALSCALE:GREGORIAN",
         f"X-WR-CALNAME:ITECH 1 - {label}", f"X-WR-TIMEZONE:{TZ}", "REFRESH-INTERVAL;VALUE=DURATION:PT1H", "X-PUBLISHED-TTL:PT1H",
         VTZ]
    n = matched = need = 0
    for ev in events:
        if group not in ev["groups"]:
            continue
        gs = ev["subs"][group]
        if sub and sub not in gs:
            continue
        lines = [l.strip() for l in ev["text"].split("\n") if l.strip()]
        title = re.sub(r"\s+", " ", lines[0]) if lines else "Cours"
        desc = [re.sub(r"\s+", " ", l) for l in lines[1:]]
        if not sub and len(gs) == 1:
            desc.append(f"Sous-groupe {next(iter(gs))} uniquement")
        rooms = find_rooms(ev, group, rooms_by_date, sub)
        loc = ""
        if not ev["allday"]:
            if rooms:
                loc = " / ".join(r["room"] for r in rooms)
                if len(rooms) > 1:
                    desc.append("Salles : " + " ; ".join(f"{r['room']}" + (f" ({r['teacher']})" if r["teacher"] else "")
                                                         for r in rooms))
            if not norm(title).startswith("e learning") and not re.match(r"\d{1,2}h", title):
                need += 1; matched += bool(rooms)
        key = f"{label}|{ev['date']}|{ev.get('start','allday')}|{norm(title)}"
        uid = hashlib.sha1(key.encode()).hexdigest()[:20] + "@planning-itech"
        L += ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{stamp}", f"SUMMARY:{esc(title)}"]
        if ev["allday"]:
            L += [f"DTSTART;VALUE=DATE:{ev['date']:%Y%m%d}", f"DTEND;VALUE=DATE:{ev['end_date']:%Y%m%d}",
                  "TRANSP:TRANSPARENT"]
        else:
            s, e = ev["start"], ev["end"]
            L += [f"DTSTART;TZID={TZ}:{ev['date']:%Y%m%d}T{s // 60:02d}{s % 60:02d}00",
                  f"DTEND;TZID={TZ}:{ev['date']:%Y%m%d}T{e // 60:02d}{e % 60:02d}00"]
        if loc:
            L.append(f"LOCATION:{esc(loc)}")
        if desc:
            L.append(f"DESCRIPTION:{esc(chr(10).join(desc))}")
        L.append("END:VEVENT"); n += 1
    L.append("END:VCALENDAR")
    return "\r\n".join(fold(l) if not l.startswith("BEGIN:VTIMEZONE") else l for l in L) + "\r\n", n, matched, need


def download(sheet_id, dest):
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        data = urllib.request.urlopen(req, timeout=120).read()
    except Exception as e:
        raise SystemExit(f"Téléchargement impossible ({url}) : {e}")
    if not data.startswith(b"PK"):
        raise SystemExit(f"Téléchargement impossible ({url}) : le Google Sheet n'est pas public ?")
    with open(dest, "wb") as f:
        f.write(data)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("planning", nargs="?"); ap.add_argument("salles", nargs="?")
    ap.add_argument("--download", action="store_true", help="télécharge les Google Sheets publics")
    ap.add_argument("--groupe", choices=list(GROUP_COLS)); ap.add_argument("--tous", action="store_true")
    ap.add_argument("--sous-groupe", choices=["A", "B"], dest="sub")
    ap.add_argument("-o", "--out"); ap.add_argument("--outdir", default=".")
    a = ap.parse_args()
    if a.download:
        a.planning, a.salles = "planning.xlsx", "salles.xlsx"
        download(PLANNING_ID, a.planning); download(SALLES_ID, a.salles)
    if not (a.planning and a.salles):
        ap.error("donne PLANNING.xlsx et SALLES.xlsx, ou utilise --download")
    events = read_planning(a.planning)
    if len(events) < MIN_EVENTS:
        raise SystemExit(f"Seulement {len(events)} événements lus : structure du fichier changée ? Rien n'est écrasé.")
    by_date = defaultdict(list)
    for r in read_rooms(a.salles):
        by_date[r["date"]].append(r)
    os.makedirs(a.outdir, exist_ok=True)
    targets = ([(g, sb) for g in GROUP_COLS for sb in (None, "A", "B")] if a.tous
               else [(a.groupe or "G1", a.sub)])
    for g, sb in targets:
        ics, n, m, need = build_ics(events, g, by_date, sb)
        out = a.out if (a.out and not a.tous) else os.path.join(a.outdir, f"itech1_{g}{sb or ''}.ics")
        with open(out, "w", encoding="utf-8", newline="") as f:
            f.write(ics)
        print(f"{out} : {n} événements, salle trouvée pour {m}/{need} cours")
    with open(os.path.join(a.outdir, "last_update.txt"), "w") as f:      # garde le dépôt "actif" pour GitHub
        f.write(dt.date.today().isoformat() + "\n")


if __name__ == "__main__":
    main()

